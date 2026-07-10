# office.py
# Office Documents Loader & Processor cho SciOS Cognitive Core

import docx
import openpyxl
import pptx
from typing import Dict, List

class OfficeProcessor:
    """
    OfficeProcessor chịu trách nhiệm:
    - Nạp và xử lý Word (.docx)
    - Nạp và xử lý Excel (.xlsx)
    - Nạp và xử lý PowerPoint (.pptx)
    """

    def __init__(self):
        pass

    # --- Word ---
    def load_word(self, file_path: str) -> List[str]:
        """Trích xuất text từ file Word."""
        doc = docx.Document(file_path)
        return [para.text for para in doc.paragraphs if para.text.strip()]

    # --- Excel ---
    def load_excel(self, file_path: str) -> Dict:
        """Trích xuất dữ liệu từ file Excel."""
        wb = openpyxl.load_workbook(file_path)
        data = {}
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
            data[sheet] = rows
        return data

    # --- PowerPoint ---
    def load_ppt(self, file_path: str) -> List[str]:
        """Trích xuất text từ file PowerPoint."""
        prs = pptx.Presentation(file_path)
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return texts

    def to_dict(self, file_path: str, file_type: str) -> Dict:
        """Chuẩn hóa output dưới dạng dict."""
        if file_type == "word":
            return {"type": "word", "content": self.load_word(file_path)}
        elif file_type == "excel":
            return {"type": "excel", "content": self.load_excel(file_path)}
        elif file_type == "ppt":
            return {"type": "ppt", "content": self.load_ppt(file_path)}
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
